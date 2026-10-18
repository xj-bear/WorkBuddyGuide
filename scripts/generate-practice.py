"""生成不含真实个人信息的教学材料；只向 practice 目录写入模拟文件。"""
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1] / 'docs/public/practice'
FILES = ROOT / 'files'


def text(name, value):
    target = FILES / name
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(value.strip() + '\n', encoding='utf-8')


def doc(name, title, paragraphs):
    target = FILES / name
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(title, 0)
    for line in paragraphs:
        document.add_paragraph(line)
    document.save(target)


def workbook(name, rows):
    target = FILES / name
    target.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = '销售明细'
    sheet.append(['订单号', '日期', '区域', '品类', '销售额', '成本'])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='126B41')
    sheet.freeze_panes = 'A2'
    for column in 'ABCDEF':
        sheet.column_dimensions[column].width = 18
    book.save(target)


def make():
    text('README.md', '''# WorkBuddy 手把手练习包
所有人物、公司、业务数据、健康与生活记录均为虚构教学材料，不代表真实情况。
解压后请选择 workbuddy-practice 作为 WorkBuddy 工作区：input 是输入，output 保存产物。
不要选择整个桌面或个人文档目录。首次演练保留 input 原件，只在 output 写结果。
六张区域销售表合计 18 行，含 1 条完全相同的跨文件重复订单。去重后 17 行，销售额合计 3250，成本 1795。
三张模拟发票合计 400.00；只用于识别和核对练习，不能用于报销。
company-facts、home-products、health-report 和生活记录均为虚构；真实研究、购买和就医另核对真实来源。
练习包不含任何预先生成的 AI 结果。每课步骤会告诉你应该生成和检查什么。
''')
    text('input/inbox/brief.txt', '虚构项目简报：本周完成 FAQ 初稿，预算待确认。')
    text('input/inbox/agenda.txt', '虚构会议议程：确认帮助中心需要更新的三个问题。')
    text('input/inbox/notes-copy.txt', '虚构会议记录副本：下周先试用，再收集反馈。')
    text('input/meeting.txt', '''教学用虚构会议转写。会议：帮助中心改版；日期：2026-10-15。
09:00 主持人：今天只讨论 FAQ 改版，不讨论发布预算。
09:03 张敏：我在 10 月 20 日前整理 10 个常见问题，交付 FAQ 初稿。
09:07 李航：我会在 10 月 22 日前完成页面原型，先给张敏看。
09:12 主持人：决定先做内部试用，不承诺正式上线日期。
09:15 陈悦：数据埋点可以看看，但今天没有确定由谁负责，也没有确定截止时间。
09:18 主持人：预算下次再讨论。本次不安排对外发送通知。
''')
    text('input/weekly-notes.md', '''# 教学用虚构工作记录
时间范围：2026-10-12 至 2026-10-16。
- 10-12：张敏完成 10 个 FAQ 问题盘点（证据：本条记录）；原型仍在进行中。
- 10-14：李航完成帮助中心线框图，已与张敏讨论；验收结果待确认。
- 10-15：决定内部试用；正式上线日期和预算未决定。
- 10-16：收集到 3 条同事反馈；下一周整理修改清单，负责人未定。
''')
    text('input/ima-exported-notes.md', '''# 虚构收藏导出，用于本地替代练习
[N1] 2026-10-12，虚构团队笔记《先给大纲》：长文先列提纲，再逐轮确认。
[N2] 2026-10-13，虚构实践记录《保留原件》：文件整理先生成清单，不直接移动。
[N3] 2026-10-14，虚构复盘《核对依据》：生成的数字和结论要能回到原文。
以上三条均为教学自编，不是实际微信收藏或真实文章。
''')
    text('input/facts.md', '''# 教学用虚构项目事实
[F1] 产品名称：小绿帮助中心；用于内部员工查询 FAQ。
[F2] 当前覆盖 10 个常见问题；计划先内部试用。
[F3] 不提供对外客服承诺；预算、上线时间尚未确定。
[F4] 本页按钮只需跳转到页面内的 FAQ 区域，不收集用户信息。
''')
    text('input/article.md', '''---
title: 先生成文件清单，再考虑整理
cover: input/cover.png
---
# 先生成文件清单，再考虑整理
工作文件杂乱时，先盘点名称、类型和路径，比直接移动更容易核对。
## 一个小例子
这是一篇虚构教学文章。练习目录只有三份文本副本；先生成清单，再检查每个文件是否还在。
## 可以带走的方法
1. 说明读取范围。
2. 先看整理建议。
3. 人工确认后，再决定是否执行。
''')
    text('input/brand-guide.md', '# 练习用视觉规范\n主色 #126B41，浅绿色 #D9FBE3，背景白色。正文使用系统无衬线字体。每页一个结论，保留来源。')
    doc('input/project-brief.docx', '虚构项目汇报资料', [
        '项目：小绿帮助中心。目标：减少内部 FAQ 重复查找。',
        '当前已有 10 个常见问题与线框图，先做内部试用。',
        '已有 3 条同事反馈。预算、上线日期和验收结论仍待确认。',
        '本文件是教学模拟材料；不得补造客户、营收或节省时长。'])
    doc('input/proposal-v1.docx', '虚构活动方案 v1', ['日期：2026-11-15', '预算：5000 元', '负责人：王敏', '仅安排内部交流，不安排对外宣传。'])
    doc('input/proposal-v2.docx', '虚构活动方案 v2', ['日期：2026-11-18', '预算：6000 元', '负责人：李芳', '安排内部交流，另准备一份会后反馈表。'])
    rows = []
    for i in range(1, 19):
        row = [f'DEMO-{i:03d}', f'2026-10-{(i-1)%9+1:02d}', f'区域{(i-1)//3+1}', '办公' if i%2 else '生活', 100+10*i, 60+5*i]
        if i == 16:
            row = rows[0].copy()
        rows.append(row)
    for n in range(6):
        workbook(f'input/sales/region-{n+1:02d}.xlsx', rows[n*3:n*3+3])
    workbook('input/sales.xlsx', list({row[0]: row for row in rows}.values()))
    for n, amount in enumerate([120.50, 80.00, 199.50], 1):
        target = FILES / f'input/invoices/invoice-demo-{n:02d}.pdf'
        target.parent.mkdir(parents=True, exist_ok=True)
        pdf = canvas.Canvas(str(target), invariant=1)
        for j, line in enumerate(['FICTIONAL TRAINING INVOICE - NOT VALID FOR REIMBURSEMENT', f'Invoice number: DEMO-{n:03d}', 'Date: 2026-10-01', f'Seller: Demo Vendor {n}', f'Amount: CNY {amount:.2f}', 'Buyer: Training Company', 'Tax ID: FICTIONAL-000']):
            pdf.drawString(45, 780-j*30, line)
        pdf.save()
    text('input/research/company-facts.md', '''# 示例科技：虚构教学证据包
这不是上市公司或实际财报。金额单位均为万元。
[R1] 虚构产品说明：提供内部知识检索软件，收入来自订阅和服务。
[R2] 虚构年度摘要：2024 收入500，净利润50，经营现金流40；2025 收入600，净利润55，经营现金流35。
[R3] 虚构风险记录：客户续约率、客户集中度、费用明细尚未披露。
[R4] 虚构项目记录：2025 增加研发投入；没有给出可核对的具体金额。
只能据此练习引用与不确定项，不推断股价、不评级、不用于投资。
''')
    text('input/brief-sources.md', '''# 本课信息来源
官网帮助文档（不是当天新闻）：
- 创建任务：https://www.codebuddy.cn/docs/workbuddy/Create-Task
- 技能：https://www.codebuddy.cn/docs/workbuddy/From-Beginner-to-Expert-Guide/Function-Description/Skills-Market
- 小程序附件：https://www.codebuddy.cn/docs/workbuddymini/features/Attachments-and-Skills
# 离线教学摘要（非新闻）
资料整理日期：2026-10-17。这是整理日期，不是下列官方页面的发布日期；三页发布日期均未知。
[S1] 来源标题：创建任务；来源URL见上。核对过的教学摘要：用户可在新建任务输入区用自然语言描述需求；文件类任务需要提供输入资料与预期结果。发布日期未知。
[S2] 来源标题：技能；来源URL见上。核对过的教学摘要：左侧“专家·技能·连接器”中的技能页签可查看市场和已安装项；安装前查看用途与权限。关闭后的技能在新会话中完全停止加载。发布日期未知。
[S3] 来源标题：小程序附件与技能；来源URL见上。核对过的教学摘要：输入栏左侧加号可添加文件或引用知识库内容；引用ima文件前先授权并选择具体文件。发布日期未知。
以上是为练习提供的简短资料摘要，不是产品更新公告，不证明今天发生新事件。离线时只能写材料型能力简报并注明未实时核验。
联网核验时记录实际查询时间，页面无发布日期时仍写“发布日期未知”。不得写成“今日新发布”。
''')
    text('input/automation/task-spec.md', '''# 同一输入重复试跑：虚构工作摘要
输入只读 input/weekly-notes.md；规则版本 v1。
先计算输入文件 SHA-256，输出 output/reliable/weekly-digest.md。
状态保存在 output/reliable/run-state.json，记录输入哈希、规则版本、状态和输出路径。
相同输入、相同规则且输出仍存在时，重跑应标记 skipped，不追加重复正文。
输入或旧产物缺失时先说明原因，不把旧结果冒充新产物；写入只在 output/reliable/。
不要连接外部系统，不发送消息。
''')
    text('input/life/travel.txt', '教学用虚构偏好：2名成人，上海出发，泉州3天，人文和小吃，预算5000元。出行日期：[填写未来日期]；住宿和交通尚未预订。步行上限15000步/日；饮食禁忌待补充。')
    text('input/life/health-report.txt', '虚构资料整理练习，不是医学报告。\n条目A：数值100，示例参考区间80–120（无医学单位）。条目B：未提供。\n记录：10-12晚间记录身体感受；10-14计划咨询医生，具体症状待本人补充。\n只练习逐字摘录、缺失项和就诊问题，不诊断、不解释正常与异常。')
    text('input/life/meals.txt', '虚构偏好：2人，5天晚餐预算300元，有电饭锅和炒锅，每餐准备30分钟；喜欢米饭、蔬菜、鸡肉。过敏信息：待本人确认，确认前不生成最终采购单。没有提供医疗饮食需求。')
    text('input/life/culture-example.txt', '虚构文化写作练习：人物小林喜欢春天与植物，希望用传统文化中的“木”意象写一段自我反思。不提供出生年月日或地点，不排盘、不预测未来，不据此做现实决定。')
    text('input/life/emotion-notes.txt', '虚构日常记录：周一赶项目有点紧张，散步后缓解；周三睡眠不足，开会容易分心；周五与朋友聊天后心情平稳。只练习整理事件与感受，不做医学或心理诊断。')
    text('input/life/trip-expenses.csv', '日期,类别,项目,金额,币种\n2026-10-01,交通,虚构车票,180,CNY\n2026-10-01,住宿,虚构民宿,240,CNY\n2026-10-01,餐饮,午餐,45,CNY\n2026-10-02,门票,虚构门票,60,CNY\n2026-10-02,餐饮,午餐,35,CNY\n2026-10-03,交通,返程接驳,90,CNY')
    text('input/life/home-products.csv', '型号,价格,清洁结构,避障,耗材年费,隐私,保修\n虚构A,2000,滚刷,激光,200,本地数据,1年\n虚构B,2500,双刷,视觉,300,云端策略未知,2年\n虚构C,1800,滚刷,未知,未知,未知,1年\n')
    cover = Image.new('RGB', (640, 320), '#D9FBE3')
    draw = ImageDraw.Draw(cover)
    draw.rectangle((64, 64, 576, 256), outline='#126B41', width=8)
    draw.rectangle((104, 110, 280, 130), fill='#126B41')
    draw.rectangle((104, 160, 480, 180), fill='#126B41')
    cover.save(FILES / 'input/cover.png')
    text('output/.keep', '本目录用于保存你在教程中生成的结果。')
    entries = sorted(p for p in FILES.rglob('*') if p.is_file())
    manifest = {'fictional': True, 'sales': {'rows': 18, 'uniqueOrders': 17, 'revenueAfterDedup': 3250, 'costAfterDedup': 1795}, 'invoices': {'count': 3, 'total': 400}, 'files': [{'path': str(p.relative_to(FILES)), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in entries]}
    (ROOT / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    with ZipFile(ROOT / 'workbuddy-practice.zip', 'w', ZIP_DEFLATED) as archive:
        for item in entries:
            archive.write(item, 'workbuddy-practice/' + item.relative_to(FILES).as_posix())
    print(f'Generated {len(entries)} fictional practice files and ZIP.')


if __name__ == '__main__':
    make()
